"""Generate file-format test cases for RAGFlow. Each file hides one unique code; a search for that
code must return the right chunk. Writes the files plus cases.json {file: {"group", "code", "expect"}}.

Usage: python format_cases.py <out dir>   (needs openpyxl, python-docx, python-pptx, reportlab, pillow)
"""
import csv, io, json, sys, wave, struct, zipfile
from pathlib import Path

out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)
cases = {}


def add(name, group, code, data, expect=None):
    p = out / name
    p.write_bytes(data if isinstance(data, bytes) else data.encode())
    cases[name] = {"group": group, "code": code, "expect": expect or code}


# 1. Extension sweep: plain-text style formats
for ext, body in {
    "txt": "Plain text note. Code {c}.", "py": "# module\nCODE = '{c}'\n", "js": "const code = '{c}';\n",
    "sql": "-- table\nSELECT '{c}' AS code;\n", "go": "package main\n// code {c}\n", "md": "# Note\nCode {c}\n",
    "mdx": "# Note\n\nCode {c}\n", "html": "<html><body><p>Code {c}</p></body></html>", "htm": "<p>Code {c}</p>",
    "json": '{{"code": "{c}"}}', "jsonl": '{{"code": "{c}"}}\n{{"other": 1}}\n',
    "eml": "From: a@example.com\nTo: b@example.com\nSubject: note\n\nCode {c}\n",
    "tsv": "name\tcode\nrow\t{c}\n", "log": "2026-10-10 10:00:00 INFO code {c}\n", "xml": "<root><code>{c}</code></root>",
    "yaml": "code: {c}\n", "yml": "code: {c}\n", "ini": "[main]\ncode={c}\n", "toml": 'code = "{c}"\n', "cfg": "code={c}\n",
    "rtf": r"{{\rtf1\ansi Code {c}\par}}",
    "arxml": '<?xml version="1.0"?><AUTOSAR><AR-PACKAGES><AR-PACKAGE><SHORT-NAME>{c}</SHORT-NAME></AR-PACKAGE></AR-PACKAGES></AUTOSAR>',
    "dbc": 'VERSION ""\nBO_ 100 EngineData: 8 ECU1\n SG_ {c} : 0|8@1+ (1,0) [0|255] "" ECU2\n',
}.items():
    c = f"FMT-{ext.upper()}-4821"
    add(f"sweep.{ext}", "sweep", c, body.format(c=c))

c = "FMT-CSV-4821"; add("sweep.csv", "sweep", c, f"name,code\nrow,{c}\n")
c = "FMT-ZIP-4821"
b = io.BytesIO(); zipfile.ZipFile(b, "w").writestr("inner.txt", f"Code {c}"); add("sweep.zip", "sweep", c, b.getvalue())
for ext, mime in [("odt", "application/vnd.oasis.opendocument.text"), ("ods", "application/vnd.oasis.opendocument.spreadsheet")]:
    c = f"FMT-{ext.upper()}-4821"; b = io.BytesIO(); z = zipfile.ZipFile(b, "w")
    z.writestr("mimetype", mime); z.writestr("content.xml", f'<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><text:p>Code {c}</text:p></office:body></office:document-content>')
    z.close(); add(f"sweep.{ext}", "sweep", c, b.getvalue())
c = "FMT-EPUB-4821"; b = io.BytesIO(); z = zipfile.ZipFile(b, "w")
z.writestr("mimetype", "application/epub+zip")
z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="c.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
z.writestr("c.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">x</dc:identifier><dc:title>T</dc:title><dc:language>en</dc:language></metadata><manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c1"/></spine></package>')
z.writestr("c1.xhtml", f'<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><body><p>Code {c}</p></body></html>')
z.close(); add("sweep.epub", "sweep", c, b.getvalue())

import docx, openpyxl, pptx
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont
c = "FMT-DOCX-4821"; d = docx.Document(); d.add_paragraph(f"Code {c}"); b = io.BytesIO(); d.save(b); add("sweep.docx", "sweep", c, b.getvalue())
c = "FMT-PPTX-4821"; p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Note"; s.placeholders[1].text = f"Code {c}"
b = io.BytesIO(); p.save(b); add("sweep.pptx", "sweep", c, b.getvalue())
c = "FMT-XLSX-4821"; wb = openpyxl.Workbook(); wb.active.append(["name", "code"]); wb.active.append(["row", c]); b = io.BytesIO(); wb.save(b); add("sweep.xlsx", "sweep", c, b.getvalue())
c = "FMT-PDF-4821"; b = io.BytesIO(); cv = canvas.Canvas(b); cv.drawString(72, 750, f"Code {c}"); cv.save(); add("sweep.pdf", "sweep", c, b.getvalue())
try:
    font = ImageFont.truetype("DejaVuSans.ttf", 44)
except OSError:
    font = ImageFont.load_default()
for ext, fmt in [("png", "PNG"), ("jpg", "JPEG")]:
    c = f"FMT-{ext.upper()}-4821"; im = Image.new("RGB", (900, 160), "white"); ImageDraw.Draw(im).text((20, 50), f"Code {c}", fill="black", font=font)
    b = io.BytesIO(); im.save(b, fmt); add(f"sweep.{ext}", "sweep", c, b.getvalue())
c = "FMT-SVG-4821"; add("sweep.svg", "sweep", c, f'<svg xmlns="http://www.w3.org/2000/svg" width="400" height="80"><text x="10" y="40">Code {c}</text></svg>')
b = io.BytesIO(); w = wave.open(b, "wb"); w.setnchannels(1); w.setsampwidth(2); w.setframerate(8000)
w.writeframes(b"".join(struct.pack("<h", int(8000 * ((i // 20) % 2 * 2 - 1))) for i in range(8000))); w.close()
add("sweep.wav", "sweep", "FMT-WAV-4821", b.getvalue())  # a tone: no words, so only acceptance/parse status is meaningful

# 2. CSV deep: target row code ROW-7777-K
T = "ROW-7777-K"
def rows(n, cols=4):
    r = [[f"item{i}"] + [f"v{i}_{j}" for j in range(1, cols)] for i in range(n)]
    r[n // 2][1] = T
    return r
def to_csv(data, delim=",", header=True, cols=4):
    s = io.StringIO(); wr = csv.writer(s, delimiter=delim)
    if header: wr.writerow(["name"] + [f"col{j}" for j in range(1, cols)])
    wr.writerows(data); return s.getvalue()
add("csv_1000rows.csv", "csv", T, to_csv(rows(1000)))
add("csv_60cols.csv", "csv", T, to_csv(rows(50, 60), cols=60))
q = rows(30); q[15][2] = 'He said "yes, no", then, later'; q[15][3] = "line one\nline two"
add("csv_quoted.csv", "csv", T, to_csv(q))
add("csv_semicolon.csv", "csv", T, to_csv(rows(30), delim=";"))
add("csv_bom.csv", "csv", T, "﻿" + to_csv(rows(30)))
add("csv_latin1.csv", "csv", T, ("name,col1,col2,col3\nCafé Müller," + T + ",Größe,Ñandú\n" + to_csv(rows(20), header=False)).encode("latin-1"))
add("csv_noheader.csv", "csv", T, to_csv(rows(30), header=False))
e = rows(30); e[15][2] = ""; e[15][3] = ""; e[3] = ["", "", "", ""]
add("csv_empty_cells.csv", "csv", T, to_csv(e))

# 3. TSV workarounds
tsv = to_csv(rows(30), delim="\t")
add("tsv_data.tsv", "tsv", T, tsv); add("tsv_as.csv", "tsv", T, tsv); add("tsv_as.txt", "tsv", T, tsv)

# 4. Excel deep
def save(wb, name, code, expect=None):
    b = io.BytesIO(); wb.save(b); add(name, "xlsx", code, b.getvalue(), expect)
wb = openpyxl.Workbook(); wb.active.title = "Summary"; wb.active.append(["a", "b"])
wb.create_sheet("Data2").append(["x", "y"]); s3 = wb.create_sheet("Data3"); s3.append(["name", "code"]); s3.append(["row", T])
save(wb, "xlsx_3sheets.xlsx", T)
wb = openpyxl.Workbook(); ws = wb.active; ws.append(["Group", "Name", "Code"]); ws.append(["Block A", "first", T]); ws.append(["", "second", "other"])
ws.merge_cells("A2:A3"); save(wb, "xlsx_merged.xlsx", T)
wb = openpyxl.Workbook(); ws = wb.active; ws.append(["Label", "Base", "Doubled"]); ws.append(["FORMULA-ROW-55", 4321, "=B2*2"])
# openpyxl writes a formula with no cached result; Excel always stores one, so add <v>8642</v> by hand.
b = io.BytesIO(); wb.save(b); zin = zipfile.ZipFile(b); b2 = io.BytesIO(); zout = zipfile.ZipFile(b2, "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == "xl/worksheets/sheet1.xml":
        data = data.replace(b"<f>B2*2</f><v></v>", b"<f>B2*2</f><v>8642</v>").replace(b"<f>B2*2</f></c>", b"<f>B2*2</f><v>8642</v></c>")
        assert b"<v>8642</v>" in data, data
    zout.writestr(it, data)
zout.close(); add("xlsx_formula.xlsx", "xlsx", "FORMULA-ROW-55", b2.getvalue(), expect="8642")  # passes only if the value is indexed
wb = openpyxl.Workbook(); ws = wb.active; ws.append(["name", "col1", "col2", "col3"])
for r in rows(1000): ws.append(r)
save(wb, "xlsx_1000rows.xlsx", T)

(out / "cases.json").write_text(json.dumps(cases, indent=1))
print(f"{len(cases)} cases written to {out}")
