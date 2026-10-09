"""Generate one small file per format, each hiding a unique fact, for the RAGFlow feature matrix.

Usage: python make_samples.py <out dir>   (needs python-docx, openpyxl, python-pptx, reportlab, pillow)
Writes the files plus facts.json: {filename: [question, [expected substrings]]}.
Facts are made up so the model cannot know them without retrieval.
"""
import csv, json, sys
from pathlib import Path

out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)
facts = {}

def fact(name, q, exp):
    facts[name] = [q, exp]

# Plain text family
(out / "notes.txt").write_text("Project Heron kickoff notes.\nThe Heron test bench is located in building K7, room 214.\n")
fact("notes.txt", "Where is the Heron test bench located?", ["K7", "214"])
(out / "readme.md").write_text("# Kestrel firmware\n\n## Release\nKestrel firmware 3.9.2 was released on 14 March 2026 by the Lisbon team.\n")
fact("readme.md", "Which team released Kestrel firmware 3.9.2?", ["Lisbon"])
(out / "page.html").write_text("<html><body><h1>Osprey CAN gateway</h1><p>The Osprey gateway bus speed is 666 kbit/s.</p></body></html>")
fact("page.html", "What is the Osprey gateway bus speed?", ["666"])
(out / "config.json").write_text(json.dumps({"product": "Merlin ECU", "watchdog_timeout_ms": 1375, "owner": "team Fjord"}))
fact("config.json", "What is the Merlin ECU watchdog timeout?", ["1375"])
(out / "mail.eml").write_text("From: anna@example.com\nTo: team@example.com\nSubject: Harrier review\n\n"
                              "The Harrier design review moved to Thursday at 15:40.\n")
fact("mail.eml", "When is the Harrier design review?", ["15:40"])

# Tables
rows = [["ECU", "Supplier", "Flash size (KB)"], ["Puffin", "Nordtek", "4096"], ["Gannet", "Velmar", "2048"], ["Tern", "Orsk", "1536"]]
with open(out / "ecus.csv", "w", newline="") as f:
    csv.writer(f).writerows(rows)
fact("ecus.csv", "Which supplier makes the Gannet ECU?", ["Velmar"])
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Budget"
for r in [["Workstream", "Budget (EUR)"], ["Cormorant validation", 48250], ["Shag integration", 17900]]:
    ws.append(r)
wb.save(out / "budget.xlsx")
fact("budget.xlsx", "What is the budget for Cormorant validation?", ["48250", "48,250", "48 250"])
(out / "signals.tsv").write_text("signal\tcycle_ms\nWheelSpeedFL\t10\nPloverStatus\t250\n")
fact("signals.tsv", "What is the cycle time of PloverStatus?", ["250"])
(out / "run.log").write_text("2026-10-01 10:00:01 INFO start\n2026-10-01 10:00:07 ERROR Skua module timeout code E-7731\n")
fact("run.log", "Which error code did the Skua module report?", ["E-7731"])

# Office
import docx
d = docx.Document(); d.add_heading("Albatross requirements", 1)
d.add_paragraph("REQ-112: The Albatross controller shall boot within 180 ms.")
t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Variant"; t.cell(0, 1).text = "Voltage"
t.cell(1, 0).text = "Albatross-HV"; t.cell(1, 1).text = "800 V"
d.save(out / "requirements.docx")
fact("requirements.docx", "Within how many milliseconds must the Albatross controller boot?", ["180"])
import pptx
from pptx.util import Inches
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "Petrel roadmap"; s.placeholders[1].text = "Petrel start of production: week 37 of 2027"
p.save(out / "roadmap.pptx")
fact("roadmap.pptx", "When is the Petrel start of production?", ["37"])

# PDF with a table (answer only in a table cell)
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, Spacer
from reportlab.lib.styles import getSampleStyleSheet
st = getSampleStyleSheet()
SimpleDocTemplate(str(out / "limits.pdf"), pagesize=A4).build([
    Paragraph("Guillemot thermal limits", st["Title"]),
    Paragraph("The table below lists the qualified operating limits.", st["Normal"]), Spacer(1, 12),
    Table([["Component", "Min (C)", "Max (C)"], ["Guillemot MCU", "-40", "125"], ["Guillemot PMIC", "-40", "105"]])])
fact("limits.pdf", "What is the maximum temperature for the Guillemot PMIC?", ["105"])

# Image with text (needs OCR)
from PIL import Image, ImageDraw, ImageFont
img = Image.new("RGB", (900, 200), "white"); dr = ImageDraw.Draw(img)
try:
    font = ImageFont.truetype("DejaVuSans.ttf", 40)
except OSError:
    font = ImageFont.load_default()
dr.text((20, 70), "Fulmar part number: FX-2291-B", fill="black", font=font)
img.save(out / "label.png")
fact("label.png", "What is the Fulmar part number?", ["FX-2291", "2291"])

(out / "facts.json").write_text(json.dumps(facts, indent=1))
print(f"wrote {len(facts)} samples to {out}")
